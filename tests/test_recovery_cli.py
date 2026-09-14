"""The public recovery workflow exports data without overwriting manuscripts."""

import hashlib
import json
import errno
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from citebind.__main__ import _write_new_output, _write_stdout, main


def manuscript(tmp_path, *, broken=False):
    document = Document()
    paragraph = document.add_paragraph("A cited claim ")
    field = OxmlElement("w:fldSimple")
    data = {"citationID": "example", "citationItems": [{
        "id": 12, "uris": ["http://zotero.org/users/1/items/ABC123"],
        "locator": "12", "label": "page", "prefix": "see ",
        "itemData": {"id": 12, "type": "book", "title": "Synthetic book",
                     "author": [{"family": "Smith", "given": "Alex"}],
                     "issued": {"date-parts": [[2020]]}},
    }]}
    field.set(qn("w:instr"), " ADDIN ZOTERO_ITEM CSL_CITATION " + (
        "{malformed" if broken else json.dumps(data)))
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "(Smith, 2020, p. 12)"
    run.append(text); field.append(run); paragraph._p.append(field)
    path = tmp_path / "manuscript.docx"
    document.save(path)
    return path


def test_cli_default_report_keeps_field_data_and_does_not_write_source(tmp_path, capsys):
    source = manuscript(tmp_path)
    original = source.read_bytes()
    assert main(["recover-references", str(source)]) == 0
    report = json.loads(capsys.readouterr().out)
    assert report["metadata_status"] == "recovered_not_verified"
    assert report["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert report["records"][0]["csl"]["type"] == "book"
    assert '"locator": "12"' in report["citations"][0]["instruction"]
    assert source.read_bytes() == original


def test_cli_creates_ris_and_discloses_unverified_status(tmp_path, capsys):
    source = manuscript(tmp_path)
    target = tmp_path / "references.ris"
    assert main(["recover-references", str(source), "--format", "ris", "--out", str(target)]) == 0
    assert "TI  - Synthetic book" in target.read_text()
    assert "unverified" in capsys.readouterr().err


def test_cli_refuses_existing_output_without_overwriting_it(tmp_path, capsys):
    source = manuscript(tmp_path)
    target = tmp_path / "references.json"
    target.write_text("Previously saved work")
    assert main(["recover-references", str(source), "--out", str(target)]) == 2
    assert target.read_text() == "Previously saved work"
    error = capsys.readouterr().err
    assert f"error: {target}: already exists" in error
    assert ".citebind-export-" not in error


def test_cli_cannot_overwrite_input(tmp_path, capsys):
    source = manuscript(tmp_path)
    before = source.read_bytes()
    assert main(["recover-references", str(source), "--out", str(source)]) == 2
    assert source.read_bytes() == before


def test_cli_broken_citation_writes_report_but_never_partial_library(tmp_path, capsys):
    source = manuscript(tmp_path, broken=True)
    report = tmp_path / "recovery.json"
    assert main(["recover-references", str(source), "--out", str(report)]) == 1
    assert any(f["severity"] == "error" for f in json.loads(report.read_text())["findings"])
    exported = tmp_path / "references.ris"
    assert main(["recover-references", str(source), "--format", "ris", "--out", str(exported)]) == 1
    assert not exported.exists()
    assert "recovery_incomplete" in capsys.readouterr().err


def test_cli_input_problem_gets_a_named_error(tmp_path, capsys):
    source = tmp_path / "bad.docx"
    source.write_text("not a zip")
    assert main(["recover-references", str(source)]) == 2
    assert "Traceback" not in capsys.readouterr().err


def test_cli_cannot_overwrite_input_through_an_output_symlink(tmp_path, capsys):
    source = manuscript(tmp_path)
    before = source.read_bytes()
    target = tmp_path / "references.json"
    target.symlink_to(source)
    assert main(["recover-references", str(source), "--out", str(target)]) == 2
    assert source.read_bytes() == before
    assert target.is_symlink()


def test_cli_full_disk_leaves_no_partial_output(tmp_path, capsys, monkeypatch):
    source = manuscript(tmp_path)
    before = source.read_bytes()
    target = tmp_path / "references.json"

    def full_disk(_fd):
        raise OSError(errno.ENOSPC, "No space left on device")

    monkeypatch.setattr(os, "fsync", full_disk)
    assert main(["recover-references", str(source), "--out", str(target)]) == 2
    assert not target.exists()
    assert source.read_bytes() == before
    assert not list(tmp_path.glob(".citebind-export-*"))


def test_cli_empty_document_does_not_export_an_empty_library(tmp_path, capsys):
    source = tmp_path / "empty.docx"
    Document().save(source)
    target = tmp_path / "references.json"
    assert main(["recover-references", str(source), "--format", "csl-json", "--out", str(target)]) == 1
    assert not target.exists()
    assert "no_recovered_references" in capsys.readouterr().err


def test_export_to_filesystem_without_hard_links(tmp_path, monkeypatch):
    def no_links(*_args):
        raise OSError(errno.EOPNOTSUPP, "No hard links")

    monkeypatch.setattr(os, "link", no_links)
    target = tmp_path / "references.json"
    _write_new_output(target, "Unicode: 中文 αβ\n")
    assert target.read_bytes() == "Unicode: 中文 αβ\n".encode("utf-8")
    import pytest
    with pytest.raises(FileExistsError):
        _write_new_output(target, "overwrite")
    assert target.read_bytes() == "Unicode: 中文 αβ\n".encode("utf-8")


def test_failed_fallback_copy_removes_only_its_partial_output(tmp_path, monkeypatch):
    import shutil
    import pytest

    def no_links(*_args):
        raise OSError(errno.EOPNOTSUPP, "No hard links")

    def full_disk(source, destination):
        destination.write(b"partial")
        raise OSError(errno.ENOSPC, "Full disk")

    monkeypatch.setattr(os, "link", no_links)
    monkeypatch.setattr(shutil, "copyfileobj", full_disk)
    target = tmp_path / "references.json"
    with pytest.raises(OSError):
        _write_new_output(target, "complete")
    assert not target.exists()
    assert not list(tmp_path.glob(".citebind-export-*"))


def test_stdout_uses_utf8_even_with_an_ascii_console(monkeypatch):
    import io
    import sys

    buffer = io.BytesIO()
    console = io.TextIOWrapper(buffer, encoding="ascii")
    monkeypatch.setattr(sys, "stdout", console)
    _write_stdout("中文 αβ\n")
    assert buffer.getvalue() == "中文 αβ\n".encode("utf-8")
