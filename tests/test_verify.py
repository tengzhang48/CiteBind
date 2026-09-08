"""T-05: structure verifier, round-trip differ, and minimal CLI.

Every damaged fixture is built by editing the DOCX XML directly — exactly
what an editor, a tool, or a bad handoff can do. The completeness test
enforces the design rule from the dev plan: the set of finding kinds is
exhaustive by construction, and adding a kind without a fixture that proves
it fails the suite.
"""

import copy
import zipfile

import pytest
from docx import Document
from lxml import etree

from citebind.controls import insert_bibliography, insert_citation
from citebind.model import CiteBindDocument
from citebind.part import embed
from citebind.__main__ import main
from citebind.rendering import render
from citebind.verify import DiffKind, FindingKind, diff, inspect

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def q(name):
    return f"{{{W_NS}}}{name}"


PAYLOAD_DICT = {
    "schema_version": "1",
    "selected_style": "numeric",
    "references": [
        {
            "id": "R001",
            "doi": "10.1000/fixture-1",
            "title": "A Fixture Paper",
            "authors": ["Alpha Author"],
            "author_names": [{"family": "Author", "given": "Alpha"}],
            "journal": "Journal of Fixtures",
            "year": 2024,
            "metadata_source": "crossref",
            "retrieved_at": "2026-08-29T00:00:00Z",
        }
    ],
    "citation_clusters": [
        {"id": "C001", "reference_ids": ["R001"]},
    ],
}


@pytest.fixture
def clean_spike(tmp_path):
    """A well-formed CiteBind document: R001, cluster C001 cited twice, bibliography.

    "Well-formed" now includes what it could not include before Phase 3: the
    visible text is RENDERED FROM THE PAYLOAD rather than typed here. A fixture
    with hand-written citation text would be a document whose visible text
    happens to disagree with its own data, which is precisely the defect
    VISIBLE_TEXT_MISMATCH exists to catch.
    """
    path = tmp_path / "clean.docx"
    rendered = render(CiteBindDocument.from_dict(copy.deepcopy(PAYLOAD_DICT)))
    document = Document()
    document.add_paragraph("As shown by prior work ")
    document.add_paragraph("and confirmed again ")
    document.add_paragraph("Ordinary prose untouched by citations.")
    paragraphs = document.paragraphs
    insert_citation(paragraphs[0], "C001", rendered.citations["C001"])
    insert_citation(paragraphs[1], "C001", rendered.citations["C001"])
    insert_bibliography(document, rendered.bibliography)
    docx_path = tmp_path / "plain.docx"
    document.save(docx_path)
    payload = CiteBindDocument.from_dict(copy.deepcopy(PAYLOAD_DICT))
    embed(docx_path, payload, path)
    return path


def rewrite_xml(src, part, transform, dst):
    with zipfile.ZipFile(src) as zin:
        tree = etree.fromstring(zin.read(part))
        transform(tree)
        new = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == part:
                    zout.writestr(item.filename, new)
                else:
                    zout.writestr(item.filename, zin.read(item.filename))


def citebind_part(path):
    with zipfile.ZipFile(path) as source:
        for name in source.namelist():
            if name.startswith("customXml/item") and name.endswith(".xml"):
                root = etree.fromstring(zipfile.ZipFile(path).read(name))
                if root.tag == "{urn:citebind:citebind:1}document":
                    return name
    raise AssertionError("no citebind payload part found")


def drop_citebind_part(src, dst):
    part = citebind_part(src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == part:
                continue
            zout.writestr(item.filename, zin.read(item.filename))


def transform_document(src, transform, dst):
    rewrite_xml(src, "word/document.xml", transform, dst)


def transform_payload(src, transform, dst):
    rewrite_xml(src, citebind_part(src), transform, dst)


def first_citation_sdt(tree):
    for sdt in tree.iter(q("sdt")):
        tag = sdt.find(f"{q('sdtPr')}/{q('tag')}")
        if tag is not None and (tag.get(q("val")) or "").startswith("citebind:citation:"):
            return sdt
    raise AssertionError("no citation control found")


# --- damage fixtures, one per acceptance case ----------------------------------


def damage_payload_removed(clean_spike, tmp_path):
    out = tmp_path / "d1.docx"
    drop_citebind_part(clean_spike, out)
    return out


def damage_controls_stripped(clean_spike, tmp_path):
    out = tmp_path / "d2.docx"

    def strip(tree):
        for sdt in list(tree.iter(q("sdt"))):
            parent = sdt.getparent()
            if parent is not None:
                parent.remove(sdt)

    transform_document(clean_spike, strip, out)
    return out


def damage_tag_renamed(clean_spike, tmp_path):
    out = tmp_path / "d3.docx"

    def rename(tree):
        for sdt in tree.iter(q("sdt")):
            tag = sdt.find(f"{q('sdtPr')}/{q('tag')}")
            if tag is not None and (tag.get(q("val")) or "").startswith("citebind:citation:"):
                tag.set(q("val"), "some:other:tool")

    transform_document(clean_spike, rename, out)
    return out


def damage_reference_deleted(clean_spike, tmp_path):
    out = tmp_path / "d4.docx"

    def delete_reference(tree):
        for node in tree.iter("{urn:citebind:citebind:1}reference"):
            node.getparent().remove(node)

    transform_payload(clean_spike, delete_reference, out)
    return out


def damage_visible_text_edited(clean_spike, tmp_path):
    out = tmp_path / "d5.docx"

    def edit_text(tree):
        t = first_citation_sdt(tree).find(f"{q('sdtContent')}/{q('r')}/{q('t')}")
        t.text = "[7]"

    transform_document(clean_spike, edit_text, out)
    return out


# --- acceptance case 6 + the five damage cases ---------------------------------


def test_clean_document_reports_zero_findings(clean_spike):
    report = inspect(clean_spike)
    assert report.findings == []


def test_plain_docx_without_citebind_reports_zero_findings(tmp_path):
    path = tmp_path / "plain.docx"
    document = Document()
    document.add_paragraph("Nothing to see here.")
    document.save(path)
    report = inspect(path)
    assert report.findings == []
    assert report.has_payload is False


def test_payload_removed_is_detected_by_name(clean_spike, tmp_path):
    report = inspect(damage_payload_removed(clean_spike, tmp_path))
    kinds = [f.kind for f in report.findings]
    assert kinds == [FindingKind.PAYLOAD_PART_MISSING]


def test_controls_stripped_is_detected_by_name(clean_spike, tmp_path):
    report = inspect(damage_controls_stripped(clean_spike, tmp_path))
    kinds = {f.kind for f in report.findings}
    assert FindingKind.CONTROL_MISSING_FOR_CLUSTER in kinds
    assert FindingKind.BIBLIOGRAPHY_CONTROL_MISSING in kinds
    assert FindingKind.CONTROL_TAG_UNRECOGNIZED not in kinds


def test_renamed_tag_is_detected_distinctly_from_stripping(clean_spike, tmp_path):
    report = inspect(damage_tag_renamed(clean_spike, tmp_path))
    kinds = {f.kind for f in report.findings}
    assert FindingKind.CONTROL_MISSING_FOR_CLUSTER in kinds
    assert FindingKind.CONTROL_TAG_UNRECOGNIZED in kinds


def test_dangling_cluster_reference_is_detected_by_name(clean_spike, tmp_path):
    report = inspect(damage_reference_deleted(clean_spike, tmp_path))
    kinds = {f.kind for f in report.findings}
    assert FindingKind.DANGLING_CLUSTER_REFERENCE in kinds
    assert FindingKind.PAYLOAD_INVALID not in kinds


def test_hand_edited_visible_text_is_detected(clean_spike, tmp_path):
    report = inspect(damage_visible_text_edited(clean_spike, tmp_path))
    kinds = {f.kind for f in report.findings}
    assert FindingKind.CITATION_CONTROLS_INCONSISTENT in kinds


# --- report content ------------------------------------------------------------


def test_report_carries_structure(clean_spike):
    report = inspect(clean_spike)
    assert report.has_payload is True
    assert report.schema_version == "1"
    assert report.selected_style == "numeric"
    assert report.reference_ids == ["R001"]
    assert report.cluster_ids == ["C001"]
    assert sorted(report.control_tags) == [
        "citebind:bibliography",
        "citebind:citation:C001",
        "citebind:citation:C001",
    ]
    assert report.inventory["sdt"] == 3
    assert report.inventory["fldChar"] == 0
    assert report.inventory["instrText"] == 0
    assert report.inventory["fldSimple"] == 0


# --- every finding kind has a fixture that proves it (exhaustive by design) ----


def fixture_duplicate_reference_id(clean_spike, tmp_path):
    out = tmp_path / "dup_ref.docx"

    def duplicate(tree):
        refs = tree.find("{urn:citebind:citebind:1}references")
        clone = copy.deepcopy(refs[0])
        refs.append(clone)

    transform_payload(clean_spike, duplicate, out)
    return out


def fixture_duplicate_cluster_id(clean_spike, tmp_path):
    out = tmp_path / "dup_cluster.docx"

    def duplicate(tree):
        clusters = tree.find("{urn:citebind:citebind:1}citation_clusters")
        clone = copy.deepcopy(clusters[0])
        clusters.append(clone)

    transform_payload(clean_spike, duplicate, out)
    return out


def fixture_payload_invalid(clean_spike, tmp_path):
    out = tmp_path / "invalid.docx"

    def drop_style(tree):
        style = tree.find("{urn:citebind:citebind:1}selected_style")
        style.getparent().remove(style)

    transform_payload(clean_spike, drop_style, out)
    return out


def fixture_reference_uncited(clean_spike, tmp_path):
    from citebind.model import Reference

    out = tmp_path / "uncited.docx"
    document = Document(clean_spike)
    payload = CiteBindDocument.from_dict(copy.deepcopy(PAYLOAD_DICT))
    payload.references.append(
        Reference(
            id="R002",
            doi="10.1000/fixture-2",
            title="Uncited Record",
            authors=["Beta Author"],
            journal="Journal of Fixtures II",
            year=2023,
            metadata_source="crossref",
            retrieved_at="2026-08-29T00:00:00Z",
        )
    )
    fresh = tmp_path / "fresh.docx"
    document.save(fresh)
    plain = tmp_path / "plain2.docx"
    document.save(plain)
    insert_citation(document.paragraphs[0], "C001", "[1]")
    insert_bibliography(document, ["A. Author. A Fixture Paper. Journal of Fixtures, 2024."])
    embed(plain, payload, out)
    return out


def fixture_control_tag_without_cluster(clean_spike, tmp_path):
    out = tmp_path / "orphan_control.docx"

    def drop_clusters(tree):
        clusters = tree.find("{urn:citebind:citebind:1}citation_clusters")
        clusters.getparent().remove(clusters)

    transform_payload(clean_spike, drop_clusters, out)
    return out


def fixture_citation_control_empty(clean_spike, tmp_path):
    out = tmp_path / "empty_text.docx"

    def empty_text(tree):
        t = first_citation_sdt(tree).find(f"{q('sdtContent')}/{q('r')}/{q('t')}")
        t.text = ""

    transform_document(clean_spike, empty_text, out)
    return out


def fixture_bibliography_count_mismatch(clean_spike, tmp_path):
    out = tmp_path / "bib_mismatch.docx"

    def add_entry(tree):
        for sdt in tree.iter(q("sdt")):
            tag = sdt.find(f"{q('sdtPr')}/{q('tag')}")
            if tag is not None and tag.get(q("val")) == "citebind:bibliography":
                content = sdt.find(q("sdtContent"))
                p = etree.SubElement(content, q("p"))
                r = etree.SubElement(p, q("r"))
                t = etree.SubElement(r, q("t"))
                t.text = "Phantom entry that nobody cited."

    transform_document(clean_spike, add_entry, out)
    return out


def fixture_visible_text_mismatch(clean_spike, tmp_path):
    """The defect that used to pass every check.

    Both citation controls are changed to "[7]" -- consistent with EACH OTHER,
    so CITATION_CONTROLS_INCONSISTENT does not fire, and structurally perfect
    in every other respect. Before Phase 3 this document inspected clean while
    showing a citation number its payload never renders.
    """
    out = tmp_path / "text_mismatch.docx"

    def rewrite_text(tree):
        for sdt in tree.iter(q("sdt")):
            tag = sdt.find(f"{q('sdtPr')}/{q('tag')}")
            if tag is not None and tag.get(q("val")) == "citebind:citation:C001":
                t = sdt.find(f"{q('sdtContent')}/{q('r')}/{q('t')}")
                t.text = "[7]"

    transform_document(clean_spike, rewrite_text, out)
    return out


def fixture_rendering_unavailable(clean_spike, tmp_path):
    """A payload the renderer refuses: structured author names removed, so no
    style can render an author label and the visible-text check cannot run.

    The note exists so that "not checked" never reads as "checked and clean".
    """
    out = tmp_path / "unrenderable.docx"

    def drop_structured_names(tree):
        for reference in tree.iter("{urn:citebind:citebind:1}author_names"):
            reference.getparent().remove(reference)

    transform_payload(clean_spike, drop_structured_names, out)
    return out


def fixture_bibliography_control_duplicate(clean_spike, tmp_path):
    """Two bibliography controls in one document.

    Not hypothetical: step 4 of the Word spike copies a citation into another
    document, and a paste that carries a bibliography control back would leave
    exactly this. The verifier used to report it CLEAN, because the scan kept
    only the last control's entry count and no finding named the duplication.
    """
    out = tmp_path / "two_bibliographies.docx"

    def duplicate_bibliography(tree):
        for sdt in list(tree.iter(q("sdt"))):
            tag = sdt.find(f"{q('sdtPr')}/{q('tag')}")
            if tag is not None and tag.get(q("val")) == "citebind:bibliography":
                sdt.addnext(copy.deepcopy(sdt))
                break

    transform_document(clean_spike, duplicate_bibliography, out)
    return out


FIXTURE_FOR_KIND = {
    FindingKind.PAYLOAD_PART_MISSING: damage_payload_removed,
    FindingKind.PAYLOAD_INVALID: fixture_payload_invalid,
    FindingKind.DUPLICATE_REFERENCE_ID: fixture_duplicate_reference_id,
    FindingKind.DUPLICATE_CLUSTER_ID: fixture_duplicate_cluster_id,
    FindingKind.DANGLING_CLUSTER_REFERENCE: damage_reference_deleted,
    FindingKind.REFERENCE_UNCITED: fixture_reference_uncited,
    FindingKind.CONTROL_MISSING_FOR_CLUSTER: damage_controls_stripped,
    FindingKind.CONTROL_TAG_WITHOUT_CLUSTER: fixture_control_tag_without_cluster,
    FindingKind.CONTROL_TAG_UNRECOGNIZED: damage_tag_renamed,
    FindingKind.CITATION_CONTROL_EMPTY: fixture_citation_control_empty,
    FindingKind.CITATION_CONTROLS_INCONSISTENT: damage_visible_text_edited,
    FindingKind.BIBLIOGRAPHY_CONTROL_MISSING: damage_controls_stripped,
    FindingKind.BIBLIOGRAPHY_COUNT_MISMATCH: fixture_bibliography_count_mismatch,
    FindingKind.VISIBLE_TEXT_MISMATCH: fixture_visible_text_mismatch,
    FindingKind.RENDERING_UNAVAILABLE: fixture_rendering_unavailable,
    FindingKind.BIBLIOGRAPHY_CONTROL_DUPLICATE: fixture_bibliography_control_duplicate,
}


def test_every_finding_kind_is_covered_by_a_fixture():
    assert set(FIXTURE_FOR_KIND) == set(FindingKind)


@pytest.mark.parametrize("kind", sorted(FindingKind, key=lambda k: k.value))
def test_each_damage_mode_is_detected_and_named(clean_spike, tmp_path, kind):
    damaged = FIXTURE_FOR_KIND[kind](clean_spike, tmp_path)
    report = inspect(damaged)
    assert kind in [f.kind for f in report.findings], (
        f"{kind.value} not detected; got {[f.kind.value for f in report.findings]}"
    )


def test_findings_carry_severity(clean_spike, tmp_path):
    report = inspect(damage_reference_deleted(clean_spike, tmp_path))
    dangling = [f for f in report.findings if f.kind == FindingKind.DANGLING_CLUSTER_REFERENCE]
    assert dangling[0].severity == "error"
    report = inspect(fixture_reference_uncited(clean_spike, tmp_path))
    uncited = [f for f in report.findings if f.kind == FindingKind.REFERENCE_UNCITED]
    assert uncited[0].severity == "warning"


# --- diff ------------------------------------------------------------------------


def test_diff_of_identical_documents_is_clean(clean_spike):
    report = diff(clean_spike, clean_spike)
    assert report.items == []


def test_diff_detects_control_text_alteration(clean_spike, tmp_path):
    damaged = damage_visible_text_edited(clean_spike, tmp_path)
    report = diff(clean_spike, damaged)
    kinds = [item.kind for item in report.items]
    assert [k.name for k in kinds] == ["CONTROL_TEXT_ALTERED"]


def test_diff_detects_payload_loss(clean_spike, tmp_path):
    damaged = damage_payload_removed(clean_spike, tmp_path)
    report = diff(clean_spike, damaged)
    assert "PAYLOAD_LOST" in [i.kind.name for i in report.items]


def test_diff_detects_introduced_fields(clean_spike, tmp_path):
    out = tmp_path / "fields.docx"

    def inject_field(tree):
        p = tree.find(f"{q('body')}/{q('p')}")
        r1 = etree.SubElement(p, q("r"))
        fld = etree.SubElement(r1, q("fldChar"))
        fld.set(q("fldCharType"), "begin")
        r2 = etree.SubElement(p, q("r"))
        instr = etree.SubElement(r2, q("instrText"))
        instr.text = " PAGE "
        r3 = etree.SubElement(p, q("r"))
        fld2 = etree.SubElement(r3, q("fldChar"))
        fld2.set(q("fldCharType"), "end")

    transform_document(clean_spike, inject_field, out)
    report = diff(clean_spike, out)
    assert "FIELD_INTRODUCED" in [i.kind.name for i in report.items]


def test_diff_detects_renamed_control(clean_spike, tmp_path):
    out = tmp_path / "renamed.docx"

    def rename(tree):
        for sdt in tree.iter(q("sdt")):
            tag = sdt.find(f"{q('sdtPr')}/{q('tag')}")
            if tag is not None and tag.get(q("val")) == "citebind:citation:C001":
                tag.set(q("val"), "citebind:citation:C009")
                return

    transform_document(clean_spike, rename, out)
    report = diff(clean_spike, out)
    names = [i.kind.name for i in report.items]
    assert "CONTROL_RENAMED" in names
    assert "CONTROL_LOST" not in names


def test_diff_detects_gained_reference(clean_spike, tmp_path):
    out = tmp_path / "more_refs.docx"
    payload = CiteBindDocument.from_dict(copy.deepcopy(PAYLOAD_DICT))
    from citebind.model import Reference

    payload.references.append(
        Reference(
            id="R002",
            doi="10.1000/fixture-2",
            title="Second Record",
            authors=["Beta Author"],
            journal="Journal of Fixtures II",
            year=2023,
            metadata_source="crossref",
            retrieved_at="2026-08-29T00:00:00Z",
        )
    )
    payload.citation_clusters.append(
        __import__("citebind.model", fromlist=["CitationCluster"]).CitationCluster(
            id="C002", reference_ids=["R002"]
        )
    )
    plain = tmp_path / "plain3.docx"
    Document(clean_spike).save(plain)
    embed(plain, payload, out)
    report = diff(clean_spike, out)
    names = [i.kind.name for i in report.items]
    assert "REFERENCE_GAINED" in names
    assert "CLUSTER_GAINED" in names


# --- CLI -------------------------------------------------------------------------


def test_cli_inspect_exit_codes(clean_spike, tmp_path, capsys):
    assert main(["inspect", str(clean_spike)]) == 0
    damaged = damage_reference_deleted(clean_spike, tmp_path)
    assert main(["inspect", str(damaged)]) == 1
    out = capsys.readouterr().out
    assert "dangling_cluster_reference" in out


def test_cli_diff_exit_codes(clean_spike, tmp_path, capsys):
    damaged = damage_visible_text_edited(clean_spike, tmp_path)
    assert main(["diff", str(clean_spike), str(clean_spike)]) == 0
    assert main(["diff", str(clean_spike), str(damaged)]) == 1
    out = capsys.readouterr().out
    assert "control_text_altered" in out


def test_cli_gives_clean_error_not_traceback_for_non_docx(tmp_path, capsys):
    bad = tmp_path / "notes.txt"
    bad.write_text("not a docx")
    assert main(["inspect", str(bad)]) == 2
    err = capsys.readouterr().err
    # message names the file and the problem; never a traceback
    assert err.startswith("error:") and "notes.txt" in err
    assert "Traceback" not in err


def test_unexpected_renderer_failure_is_a_note_not_a_traceback(clean_spike, monkeypatch):
    """inspect exists to report on documents that are ALREADY damaged, so a
    renderer that throws on a hostile payload must become a finding. A verifier
    that dies on bad input has no answer for the case it was built for."""
    import citebind.verify as verify_module

    def exploding_render(_document):
        raise ValueError("renderer went bang")

    monkeypatch.setattr(verify_module, "render", exploding_render)
    report = inspect(clean_spike)
    notes = [f for f in report.findings if f.kind == FindingKind.RENDERING_UNAVAILABLE]
    assert notes, "an exploding renderer must still produce a report"
    assert "ValueError" in notes[0].detail and "went bang" in notes[0].detail


def test_payload_round_trip_preserves_structured_author_names(tmp_path):
    """The split is only useful if it survives the document. embed/extract is
    the path every rendered citation depends on."""
    from citebind.part import extract

    document = Document()
    document.add_paragraph("Prose.")
    plain = tmp_path / "plain.docx"
    document.save(plain)
    payload = CiteBindDocument.from_dict(copy.deepcopy(PAYLOAD_DICT))
    out = tmp_path / "round_trip.docx"
    embed(plain, payload, out)
    recovered = extract(out)
    assert recovered.references[0].author_names == payload.references[0].author_names
    assert recovered.to_dict() == payload.to_dict()


def test_payload_missing_authors_is_reported_not_a_traceback(clean_spike, tmp_path):
    """REGRESSION. Deleting the <authors> element used to take inspect down
    with a bare KeyError from the schema, which is the worst possible failure
    for a tool whose entire purpose is reporting on documents that are already
    damaged. The damage must come back as a named finding."""
    out = tmp_path / "no_authors.docx"

    def drop_authors(tree):
        for node in list(tree.iter("{urn:citebind:citebind:1}authors")):
            node.getparent().remove(node)

    transform_payload(clean_spike, drop_authors, out)
    report = inspect(out)
    kinds = [f.kind for f in report.findings]
    assert FindingKind.PAYLOAD_INVALID in kinds, kinds
    assert "authors" in " ".join(f.detail for f in report.findings)


def _rewrite_payload_field(source, out, transform):
    """Edit one field inside the payload part of a copy of ``source``."""

    def apply(tree):
        transform(tree)

    transform_payload(source, apply, out)
    return out


def test_diff_reports_a_reference_whose_metadata_was_rewritten(clean_spike, tmp_path):
    """REGRESSION. diff compared reference IDs by presence only, so a handoff
    that rewrote a reference's year or title came back "no differences" — the
    id was still in the list, so nothing had "changed". Presence was standing
    in for identity: a reference is what it claims, not that its id survived.
    """
    out = tmp_path / "rewritten.docx"

    def rewrite_year(tree):
        year = tree.find(
            "{urn:citebind:citebind:1}references/"
            "{urn:citebind:citebind:1}reference/{urn:citebind:citebind:1}year"
        )
        year.text = "1999"

    _rewrite_payload_field(clean_spike, out, rewrite_year)
    report = diff(clean_spike, out)
    changed = [i for i in report.items if i.kind == DiffKind.REFERENCE_CHANGED]
    assert changed, [i.kind.value for i in report.items]
    assert "1999" in changed[0].detail


def test_diff_reports_a_reference_that_lost_a_required_field(clean_spike, tmp_path):
    """The sharper case: the document went from valid to INVALID — a required
    field deleted outright — and diff reported no differences at all."""
    out = tmp_path / "field_deleted.docx"

    def drop_authors(tree):
        for node in list(tree.iter("{urn:citebind:citebind:1}authors")):
            node.getparent().remove(node)

    _rewrite_payload_field(clean_spike, out, drop_authors)
    report = diff(clean_spike, out)
    assert DiffKind.REFERENCE_CHANGED in [i.kind for i in report.items]


def test_diff_of_a_document_with_itself_is_still_clean(clean_spike):
    """The content comparison must not invent differences."""
    assert diff(clean_spike, clean_spike).items == []


# --- diff completeness, mirroring the rule already enforced for findings -----
#
# FindingKind has had an exhaustiveness gate since T-05: a kind without a
# fixture that produces it fails the suite. DiffKind had no such gate, which
# is how REFERENCE_CHANGED and CLUSTER_CHANGED were added with nothing forcing
# a demonstration. Each entry below must be a PAIR of documents that really
# produces its kind.


def diff_pair_payload_lost(clean_spike, tmp_path):
    out = tmp_path / "d_lost.docx"
    drop_citebind_part(clean_spike, out)
    return clean_spike, out


def diff_pair_payload_gained(clean_spike, tmp_path):
    out = tmp_path / "d_gained.docx"
    drop_citebind_part(clean_spike, out)
    return out, clean_spike


def diff_pair_schema_version_changed(clean_spike, tmp_path):
    out = tmp_path / "d_schema.docx"

    def bump(tree):
        tree.set("schema_version", "2")

    transform_payload(clean_spike, bump, out)
    return clean_spike, out


def diff_pair_style_changed(clean_spike, tmp_path):
    out = tmp_path / "d_style.docx"

    def restyle(tree):
        tree.find("{urn:citebind:citebind:1}selected_style").text = "author-year"

    transform_payload(clean_spike, restyle, out)
    return clean_spike, out


def diff_pair_reference_lost(clean_spike, tmp_path):
    return clean_spike, damage_reference_deleted(clean_spike, tmp_path)


def diff_pair_reference_gained(clean_spike, tmp_path):
    return damage_reference_deleted(clean_spike, tmp_path), clean_spike


def diff_pair_cluster_lost(clean_spike, tmp_path):
    return clean_spike, fixture_control_tag_without_cluster(clean_spike, tmp_path)


def diff_pair_cluster_gained(clean_spike, tmp_path):
    return fixture_control_tag_without_cluster(clean_spike, tmp_path), clean_spike


def diff_pair_control_lost(clean_spike, tmp_path):
    return clean_spike, damage_controls_stripped(clean_spike, tmp_path)


def diff_pair_control_gained(clean_spike, tmp_path):
    return damage_controls_stripped(clean_spike, tmp_path), clean_spike


def diff_pair_control_renamed(clean_spike, tmp_path):
    out = tmp_path / "d_renamed.docx"

    def rename(tree):
        tag = first_citation_sdt(tree).find(f"{q('sdtPr')}/{q('tag')}")
        tag.set(q("val"), "citebind:citation:C999")

    transform_document(clean_spike, rename, out)
    return clean_spike, out


def diff_pair_control_text_altered(clean_spike, tmp_path):
    return clean_spike, damage_visible_text_edited(clean_spike, tmp_path)


def diff_pair_reference_changed(clean_spike, tmp_path):
    out = tmp_path / "d_ref_changed.docx"

    def rewrite(tree):
        tree.find(
            "{urn:citebind:citebind:1}references/"
            "{urn:citebind:citebind:1}reference/{urn:citebind:citebind:1}year"
        ).text = "1999"

    transform_payload(clean_spike, rewrite, out)
    return clean_spike, out


def diff_pair_cluster_changed(clean_spike, tmp_path):
    out = tmp_path / "d_cluster_changed.docx"

    def add_locator(tree):
        cluster = tree.find(
            "{urn:citebind:citebind:1}citation_clusters/"
            "{urn:citebind:citebind:1}cluster"
        )
        etree.SubElement(cluster, "{urn:citebind:citebind:1}locator").text = "12"

    transform_payload(clean_spike, add_locator, out)
    return clean_spike, out


def diff_pair_field_introduced(clean_spike, tmp_path):
    """A Word field appears where CiteBind uses content controls — what a
    reference manager's own citation would look like arriving in the file."""
    out = tmp_path / "d_field.docx"

    def add_field(tree):
        body = tree.find(q("body"))
        paragraph = etree.SubElement(body, q("p"))
        etree.SubElement(paragraph, q("fldSimple")).set(q("instr"), "CITATION Bel10")

    transform_document(clean_spike, add_field, out)
    return clean_spike, out


def diff_pair_sdt_count_changed(clean_spike, tmp_path):
    """An sdt appears that no citebind tag accounts for, so the control-tag
    bookkeeping cannot explain the change in sdt count."""
    out = tmp_path / "d_sdt_count.docx"

    def add_foreign_sdt(tree):
        body = tree.find(q("body"))
        sdt = etree.SubElement(body, q("sdt"))
        properties = etree.SubElement(sdt, q("sdtPr"))
        etree.SubElement(properties, q("tag")).set(q("val"), "someone-elses-control")
        etree.SubElement(sdt, q("sdtContent"))

    transform_document(clean_spike, add_foreign_sdt, out)
    return clean_spike, out


PAIR_FOR_DIFF_KIND = {
    DiffKind.PAYLOAD_LOST: diff_pair_payload_lost,
    DiffKind.PAYLOAD_GAINED: diff_pair_payload_gained,
    DiffKind.SCHEMA_VERSION_CHANGED: diff_pair_schema_version_changed,
    DiffKind.STYLE_CHANGED: diff_pair_style_changed,
    DiffKind.REFERENCE_LOST: diff_pair_reference_lost,
    DiffKind.REFERENCE_GAINED: diff_pair_reference_gained,
    DiffKind.CLUSTER_LOST: diff_pair_cluster_lost,
    DiffKind.CLUSTER_GAINED: diff_pair_cluster_gained,
    DiffKind.CONTROL_LOST: diff_pair_control_lost,
    DiffKind.CONTROL_GAINED: diff_pair_control_gained,
    DiffKind.CONTROL_RENAMED: diff_pair_control_renamed,
    DiffKind.CONTROL_TEXT_ALTERED: diff_pair_control_text_altered,
    DiffKind.REFERENCE_CHANGED: diff_pair_reference_changed,
    DiffKind.CLUSTER_CHANGED: diff_pair_cluster_changed,
    DiffKind.FIELD_INTRODUCED: diff_pair_field_introduced,
    DiffKind.SDT_COUNT_CHANGED: diff_pair_sdt_count_changed,
}


def test_every_diff_kind_is_covered_by_a_pair():
    assert set(PAIR_FOR_DIFF_KIND) == set(DiffKind)


@pytest.mark.parametrize("kind", sorted(DiffKind, key=lambda k: k.value))
def test_each_diff_kind_is_produced_by_its_pair(clean_spike, tmp_path, kind):
    before, after = PAIR_FOR_DIFF_KIND[kind](clean_spike, tmp_path)
    report = diff(before, after)
    assert kind in [item.kind for item in report.items], (
        f"{kind.value} not produced; got {[i.kind.value for i in report.items]}"
    )
