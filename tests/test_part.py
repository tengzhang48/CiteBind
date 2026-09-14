"""T-03: embedding and recovering the citebind/1 custom XML part in a DOCX.

The dev plan's brief said to write ``customXml/item1.xml``, but python-docx's
default template — like real Word documents — already uses that part for
Word's built-in bibliography sources (``b:Sources``). Overwriting it would be
destructive, so CiteBind takes the next free ``customXml/itemN.xml`` slot and
reuses it on re-embed. These tests pin that behavior, including the
non-negotiable guarantee that Word's own item1 part survives untouched.
"""

import re
import zipfile

import pytest
from lxml import etree
from docx import Document

from citebind.model import CiteBindDocument
from citebind.part import PayloadError, embed, extract, find_citebind_part
from test_schema import wellformed_full, wellformed_minimal_doi


@pytest.fixture
def base_docx(tmp_path):
    path = tmp_path / "base.docx"
    document = Document()
    document.add_paragraph("First paragraph of ordinary prose.")
    document.add_paragraph("Second paragraph of ordinary prose.")
    document.save(path)
    return path


def rewrite_entry(src, name, data, dst):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == name:
                zout.writestr(item.filename, data)
            else:
                zout.writestr(item.filename, zin.read(item.filename))


def payload_doc():
    return CiteBindDocument.from_dict(wellformed_full())


def citebind_part_name(path):
    with zipfile.ZipFile(path) as source:
        name = find_citebind_part(source)
    assert name is not None
    return name


def item_number(part_name):
    return int(re.match(r"^customXml/item(\d+)\.xml$", part_name).group(1))


# --- acceptance 1: the part is present after embed ----------------------------


def test_embed_adds_custom_xml_part(base_docx, tmp_path):
    out = tmp_path / "out.docx"
    used = embed(base_docx, payload_doc(), out)
    with zipfile.ZipFile(out) as source:
        names = source.namelist()
        assert find_citebind_part(source) == used
    assert used == "customXml/item2.xml"  # item1 is Word's b:Sources; we take 2
    n = item_number(used)
    assert f"customXml/itemProps{n}.xml" in names
    assert f"customXml/_rels/item{n}.xml.rels" in names
    content_types = zipfile.ZipFile(out).read("[Content_Types].xml")
    n = item_number(used)
    # The template covers .xml via Default, and itemProps1 uses the legacy
    # ...customXmlProperties+xml type; we mirror the document's own convention.
    assert b'PartName="/customXml/itemProps2.xml"' in content_types
    assert b'PartName="/customXml/item2.xml"' not in content_types
    assert b'<Default Extension="xml"' in content_types
    # The relationship belongs to the MAIN DOCUMENT part, not the package
    # root: that is where Word puts its own (word/_rels/document.xml.rels ->
    # "../customXml/item1.xml"), and a relationship graph Word does not follow
    # is a payload Word cannot see, however well our own reader finds it.
    document_rels = zipfile.ZipFile(out).read("word/_rels/document.xml.rels")
    assert b"relationships/customXml" in document_rels
    assert f"../customXml/item{n}.xml".encode() in document_rels
    package_rels = zipfile.ZipFile(out).read("_rels/.rels")
    assert f"customXml/item{n}.xml".encode() not in package_rels


def test_embed_preserves_words_own_bibliography_sources(base_docx, tmp_path):
    out = tmp_path / "out.docx"
    embed(base_docx, payload_doc(), out)
    original_item1 = zipfile.ZipFile(base_docx).read("customXml/item1.xml")
    embedded_item1 = zipfile.ZipFile(out).read("customXml/item1.xml")
    assert original_item1 == embedded_item1


# --- acceptance 2: extract returns an equal model ------------------------------


def test_extract_returns_equal_model(base_docx, tmp_path):
    out = tmp_path / "out.docx"
    embed(base_docx, payload_doc(), out)
    assert extract(out) == payload_doc()


# --- acceptance 3: python-docx still opens the output --------------------------


def test_python_docx_still_opens_output(base_docx, tmp_path):
    out = tmp_path / "out.docx"
    embed(base_docx, payload_doc(), out)
    document = Document(out)
    texts = [p.text for p in document.paragraphs]
    assert texts == [
        "First paragraph of ordinary prose.",
        "Second paragraph of ordinary prose.",
    ]


# --- acceptance 4: the load-bearing one — python-docx full rewrite -------------


def test_payload_survives_python_docx_rewrite_byte_identical(base_docx, tmp_path):
    out = tmp_path / "out.docx"
    embed(base_docx, payload_doc(), out)
    resaved = tmp_path / "resaved.docx"
    Document(out).save(resaved)
    part = citebind_part_name(out)
    with zipfile.ZipFile(out) as z1, zipfile.ZipFile(resaved) as z2:
        assert z1.read(part) == z2.read(part)
        n = item_number(part)
        assert z1.read(f"customXml/itemProps{n}.xml") == z2.read(
            f"customXml/itemProps{n}.xml"
        )
    assert extract(resaved) == payload_doc()


# --- acceptance 5: no CiteBind part means None, not an exception ---------------


def test_extract_returns_none_without_part(base_docx):
    assert extract(base_docx) is None


# --- acceptance 6: DOCTYPE in the payload is refused, by name ------------------


def test_extract_refuses_doctype(base_docx, tmp_path):
    out = tmp_path / "out.docx"
    embed(base_docx, payload_doc(), out)
    laced = tmp_path / "laced.docx"
    xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b"<!DOCTYPE document [<!ENTITY xxe \"oops\">]>"
        b'<citebind:document xmlns:citebind="urn:citebind:citebind:1" '
        b'schema_version="1"/>'
    )
    rewrite_entry(out, citebind_part_name(out), xml, laced)
    with pytest.raises(PayloadError) as exc_info:
        extract(laced)
    assert exc_info.value.code == "doctype_declared"


def test_extract_returns_none_when_slot_root_is_foreign(base_docx, tmp_path):
    # Extraction is content-addressed: an item whose root is not ours simply
    # is not a CiteBind payload. The brief's "refuse foreign root on extract"
    # scenario cannot arise under slot allocation; the DOCTYPE refusal above
    # is the loud failure path that remains. (Deliberate contract change.)
    out = tmp_path / "out.docx"
    embed(base_docx, payload_doc(), out)
    foreign = tmp_path / "foreign.docx"
    rewrite_entry(
        out,
        citebind_part_name(out),
        b'<other xmlns="urn:example:other"/>',
        foreign,
    )
    assert extract(foreign) is None


def test_embed_allocates_fresh_slot_when_payload_was_clobbered(base_docx, tmp_path):
    # A clobbered citebind slot is never overwritten (nothing foreign is ever
    # destroyed); embed allocates the next slot instead. (Deliberate contract
    # change from the brief's "refuse foreign overwrite": the only item embed
    # could clobber under the old plan was Word's own item1, and slot
    # allocation makes that impossible by construction.)
    out = tmp_path / "out.docx"
    embed(base_docx, payload_doc(), out)
    foreign = tmp_path / "foreign.docx"
    rewrite_entry(
        out,
        citebind_part_name(out),
        b'<other xmlns="urn:example:other"/>',
        foreign,
    )
    fresh = tmp_path / "fresh.docx"
    used = embed(foreign, payload_doc(), fresh)
    assert used == "customXml/item3.xml"
    assert zipfile.ZipFile(fresh).read("customXml/item2.xml") == (
        b'<other xmlns="urn:example:other"/>'
    )
    assert extract(fresh) == payload_doc()


def test_reembed_reuses_same_slot_and_updates_payload(base_docx, tmp_path):
    first = tmp_path / "first.docx"
    used_first = embed(base_docx, CiteBindDocument.from_dict(wellformed_minimal_doi()), first)
    second = tmp_path / "second.docx"
    used_second = embed(first, payload_doc(), second)
    assert used_first == used_second
    assert extract(second) == payload_doc()
    assert zipfile.ZipFile(second).namelist().count("customXml/item2.xml") == 1


# --- F4: two payload parts is a named refusal, not a coin flip -----------------


def test_two_payload_parts_are_refused_by_name(base_docx, tmp_path):
    out = tmp_path / "one.docx"
    embed(base_docx, payload_doc(), out)
    dual = tmp_path / "dual.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(dual, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item.filename, zin.read(item.filename))
        # a second citebind payload part, as if a tool had cloned the slot
        zout.writestr("customXml/item4.xml", zipfile.ZipFile(out).read(citebind_part_name(out)))
    with pytest.raises(PayloadError) as exc_info:
        extract(dual)
    assert exc_info.value.code == "multiple_payload_parts"


def test_inspect_reports_ambiguous_payload_as_finding(base_docx, tmp_path):
    from citebind.verify import FindingKind, inspect

    out = tmp_path / "one.docx"
    embed(base_docx, payload_doc(), out)
    dual = tmp_path / "dual.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(dual, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item.filename, zin.read(item.filename))
        zout.writestr("customXml/item4.xml", zipfile.ZipFile(out).read(citebind_part_name(out)))
    report = inspect(dual)
    assert FindingKind.PAYLOAD_INVALID in [f.kind for f in report.findings]
    assert any("multiple_payload_parts" in f.detail for f in report.findings)


def test_datastore_item_id_is_a_valid_ooxml_guid(tmp_path):
    """REGRESSION, and SOURCED against the spec rather than against ourselves.

    ECMA-376 types ``ds:itemID`` as ``ST_Guid``, whose pattern requires
    surrounding braces. We wrote a bare GUID, so every document CiteBind
    produced carried an attribute value Word is entitled to reject -- and the
    part python-docx's own template writes, right next to ours, has the braces.

    The cost of getting this wrong was not a crash: it was a Word repair prompt
    on the Phase 1 spike fixture, which would have been read as evidence that
    content controls do not survive Word.
    """
    st_guid = re.compile(
        r"^\{[0-9A-F]{8}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{12}\}$"
    )

    document = Document()
    document.add_paragraph("Body.")
    plain = tmp_path / "plain.docx"
    document.save(plain)
    out = tmp_path / "out.docx"
    embed(plain, CiteBindDocument.from_dict(wellformed_minimal_doi()), out)
    with zipfile.ZipFile(out) as package:
        props = [n for n in package.namelist() if "itemProps" in n]
        found = []
        for name in props:
            match = re.search(r'itemID="([^"]+)"', package.read(name).decode())
            if match:
                found.append(match.group(1))
    assert found, "no itemProps part carried an itemID"
    for value in found:
        assert st_guid.match(value), f"{value} is not a valid ST_Guid"


def test_item_props_declares_our_namespace_like_word_declares_its_own(base_docx, tmp_path):
    """Word's own custom XML part in the same package carries a schemaRefs
    naming the namespace of its data (the bibliography namespace). Ours carried
    none, making it the only custom XML part in the document that did not say
    what it holds.

    The element is optional, so this was never a conformance failure. It is
    here because the Word spike is about to interpret whatever Word does to
    this part, and an unexplained difference from Word's own shape would be a
    live candidate explanation for any failure observed.
    """
    out = tmp_path / "props.docx"
    used = embed(base_docx, CiteBindDocument.from_dict(wellformed_minimal_doi()), out)
    n = item_number(used)
    ds = "http://schemas.openxmlformats.org/officeDocument/2006/customXml"
    props = etree.fromstring(
        zipfile.ZipFile(out).read(f"customXml/itemProps{n}.xml")
    )
    refs = props.findall(f"{{{ds}}}schemaRefs/{{{ds}}}schemaRef")
    assert [r.get(f"{{{ds}}}uri") for r in refs] == ["urn:citebind:citebind:1"]


def test_embed_refuses_a_package_that_is_not_a_word_document(base_docx, tmp_path):
    """embed reads three parts to wire the payload in. A package missing one
    raised a bare KeyError from inside zipfile, three frames below the caller —
    the same unnamed-crash shape this library keeps closing elsewhere."""
    import zipfile as zf

    stripped = tmp_path / "stripped.docx"
    with zf.ZipFile(base_docx) as source, zf.ZipFile(stripped, "w") as target:
        for item in source.infolist():
            if item.filename == "word/_rels/document.xml.rels":
                continue
            target.writestr(item.filename, source.read(item.filename))
    with pytest.raises(PayloadError) as caught:
        embed(stripped, CiteBindDocument.from_dict(wellformed_minimal_doi()), tmp_path / "o.docx")
    assert caught.value.code == "not_a_word_package"
