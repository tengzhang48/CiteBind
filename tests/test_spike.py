"""T-06: the Word round-trip kit — generator, checker, and their contract.

The manual Word step cannot be run on this machine (dev plan §2). These tests
do the two things that CAN be tested here: the generator produces exactly the
promised fixture with no network, and the checker reaches the right per-probe
verdict on synthetic stand-ins for every probe outcome.

Checker design rules (review note 2026-08-29, F1/F2): files are identified by
the names the instructions promise; P1/P4/P6 each grade their own evidence
file; every row names the file it read and the fact that decided it.
"""

import copy
import socket
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from citebind.controls import find_controls, insert_citation
from citebind.part import extract
from citebind.spike import (
    BASELINE_DOI,
    BASELINE_REFERENCE,
    check_spike,
    make_spike,
)
from test_verify import q  # same W-namespace helper

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# --- the generator ---------------------------------------------------------------


def test_make_spike_produces_promised_fixture(tmp_path):
    path = make_spike(tmp_path)
    assert path.name == "spike_v1.docx"

    document = Document(path)  # opens cleanly in python-docx
    # two prose paragraphs; the bibliography lives inside a content control,
    # which python-docx's high-level API does not list (documented in T-04)
    assert len(document.paragraphs) == 2

    controls = find_controls(path)
    citations = [c for c in controls if c.kind == "citation"]
    bibliographies = [c for c in controls if c.kind == "bibliography"]
    assert len(citations) == 2
    assert all(c.cluster_id == "C001" for c in citations)
    assert all(c.text == "[1]" for c in citations)
    assert len(bibliographies) == 1

    payload = extract(path)
    assert payload is not None
    assert len(payload.references) == 1
    reference = payload.references[0]
    assert reference.doi == BASELINE_DOI
    assert reference.to_dict() == BASELINE_REFERENCE
    assert payload.citation_clusters[0].id == "C001"
    assert payload.citation_clusters[0].reference_ids == ["R001"]


def test_spike_payload_is_the_archived_crossref_record(tmp_path):
    # The hard-coded reference must be exactly the record archived from
    # Crossref (spike/reference_source_crossref.json) — never from memory.
    import json

    path = make_spike(tmp_path)
    payload = extract(path)
    archived = json.load(open(Path(__file__).parent.parent / "spike" / "reference_source_crossref.json"))
    message = archived["message"]
    assert payload.references[0].doi == message["DOI"]
    assert payload.references[0].title == message["title"][0]
    assert payload.references[0].journal == message["container-title"][0]


def test_make_spike_uses_no_network(tmp_path, monkeypatch):
    def refuse(*args, **kwargs):
        raise AssertionError("make_spike attempted a network call")

    monkeypatch.setattr(socket, "socket", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)
    make_spike(tmp_path)


def test_generated_spike_introduces_no_fields(tmp_path):
    path = make_spike(tmp_path)
    xml = zipfile.ZipFile(path).read("word/document.xml")
    tree = etree.fromstring(xml)
    for name in ("fldChar", "instrText", "fldSimple"):
        assert len(list(tree.iter(q(name)))) == 0


# --- simulated Word sessions ------------------------------------------------------


def edit_document_xml(src, dst, fn):
    with zipfile.ZipFile(src) as zin:
        tree = etree.fromstring(zin.read("word/document.xml"))
        fn(tree)
        new = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item.filename, new if item.filename == "word/document.xml" else zin.read(item.filename))


def drop_payload_parts(src, dst):
    """Remove CiteBind's payload part and its package relationship, the way a
    package that 'lost the payload' during a Word save would look."""
    import re as _re

    from citebind.part import find_citebind_part

    with zipfile.ZipFile(src) as zin:
        ours = find_citebind_part(zin)
        assert ours is not None
        n = _re.match(r"^customXml/item(\d+)\.xml$", ours).group(1)
        drop = {ours, f"customXml/itemProps{n}.xml", f"customXml/_rels/item{n}.xml.rels"}
        rels = etree.fromstring(zin.read("_rels/.rels"))
        for rel in list(rels):
            if rel.get("Target") == ours:
                rels.remove(rel)
        new_rels = etree.tostring(rels, xml_declaration=True, encoding="UTF-8")
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in drop:
                    continue
                if item.filename == "_rels/.rels":
                    zout.writestr(item.filename, new_rels)
                else:
                    zout.writestr(item.filename, zin.read(item.filename))


def first_citation_sdt(tree):
    for sdt in tree.iter(q("sdt")):
        tag = sdt.find(f"{q('sdtPr')}/{q('tag')}")
        if tag is not None and tag.get(q("val")) == "citebind:citation:C001":
            return sdt
    raise AssertionError("no citation control found")


def resave(path, name):
    """Round a file through python-docx's writer, simulating a Word re-save."""
    out = path.parent / name
    Document(path).save(out)
    return out


def make_pasted_doc(tmp_path, name="pasted.docx", with_payload=False, pristine=None):
    if with_payload:
        pasted = tmp_path / name
        copy_doc(pristine, pasted)
        return pasted
    pasted = tmp_path / name
    document = Document()
    document.add_paragraph("Pasted into a blank document ")
    insert_citation(document.paragraphs[0], "C001", "[1]")
    document.save(pasted)
    return pasted


def copy_doc(src, dst):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item.filename, zin.read(item.filename))


def build_returned_files(
    tmp_path,
    prose=True,
    paste=True,
    tracked=True,
    payload_loss_at_end=False,
    pasted_with_payload=False,
    include_step1=True,
    include_pasted=True,
    include_renamed=True,
    tamper_text=False,
):
    """Simulate one human Word session and return the four promised files."""
    pristine = make_spike(tmp_path)
    pristine_keep = tmp_path / "pristine_keep.docx"
    copy_doc(pristine, pristine_keep)  # main will overwrite spike_v1.docx below

    def session(tree):
        prose_paragraphs = tree.find(q("body")).findall(q("p"))
        if prose:
            r = etree.SubElement(prose_paragraphs[0], q("r"))
            t = etree.SubElement(r, q("t"))
            t.text = " A typed sentence. "
        if paste:
            prose_paragraphs[1].append(copy.deepcopy(first_citation_sdt(tree)))
        if tracked:
            first = first_citation_sdt(tree)
            ins = etree.Element(q("ins"))
            ins.set(q("id"), "1")
            run = etree.SubElement(ins, q("r"))
            t = etree.SubElement(run, q("t"))
            t.text = "Tracked words "
            first.addprevious(ins)
        if tamper_text:
            t = first_citation_sdt(tree).find(f"{q('sdtContent')}/{q('r')}/{q('t')}")
            t.text = "[7]"

    edited = tmp_path / "edited.docx"
    edit_document_xml(pristine, edited, session)
    main = resave(edited, "spike_v1.docx")
    if payload_loss_at_end:
        damaged = tmp_path / "main_damaged.docx"
        drop_payload_parts(main, damaged)
        main = resave(damaged, "spike_v1.docx")

    step1 = resave(pristine_keep, "step1_reopened.docx") if include_step1 else None
    pasted = (
        make_pasted_doc(tmp_path, with_payload=pasted_with_payload, pristine=pristine_keep)
        if include_pasted
        else None
    )
    renamed = resave(main, "spike_renamed.docx") if include_renamed else None

    files = [main]
    files += [f for f in (step1, pasted, renamed) if f is not None]
    return files


# --- checker verdicts -------------------------------------------------------------


def row(report, probe):
    return [r for r in report.rows if r.probe == probe][0]


def test_check_spike_all_probes_pass(tmp_path):
    files = build_returned_files(tmp_path)
    report = check_spike(files)
    for r in report.rows:
        assert r.verdict == "PASS", f"{r.probe}: {r.detail}"
    assert report.all_passed
    assert any("graded as" in line for line in report.classification)


def test_f1_step5_payload_loss_must_not_fail_p1(tmp_path):
    # The regression the review note demands: a defect introduced at step 5
    # (payload lost during the Track Changes save) must NOT fail P1 — P1 is
    # graded on its own step-1 artifact, which is intact here.
    files = build_returned_files(tmp_path, payload_loss_at_end=True)
    report = check_spike(files)
    verdicts = {r.probe: r.verdict for r in report.rows}
    assert verdicts["P1"] == "PASS", row(report, "P1").detail
    # the loss IS caught, by the probe that owns the end state:
    assert verdicts["P6"] == "FAIL"
    assert "payload missing or altered" in row(report, "P6").detail
    # P2 and P5 grade the visible layer, which survived; their rows say
    # exactly that instead of borrowing payload evidence (old P5 printed
    # FAIL above a detail claiming everything was intact)
    assert verdicts["P2"] == "PASS", row(report, "P2").detail
    assert verdicts["P5"] == "PASS"
    assert "tracked edits present" in row(report, "P5").detail
    # every failing row names the file it read and a deciding fact
    for r in report.rows:
        if r.verdict == "FAIL":
            assert "deciding fact" in r.detail or "was not returned" in r.detail


def test_f2_paste_retaining_payload_is_still_classified_as_paste(tmp_path):
    # The regression the review note demands: when a clipboard paste carries
    # the payload too, files must keep the roles their names promise. The old
    # content-based classifier graded pasted.docx as the save-as copy and
    # asserted the library "does not travel (expected)" — the exact inverse.
    files = build_returned_files(tmp_path, pasted_with_payload=True)
    report = check_spike(files)
    p4 = row(report, "P4")
    p6 = row(report, "P6")
    assert p4.file_read.endswith("pasted.docx")
    assert p6.file_read.endswith("spike_renamed.docx")
    # the payload's presence is reported as a fact, with no "expected" label
    assert "payload part present" in p4.detail
    assert "expected" not in p4.detail
    assert p6.verdict == "PASS"


def test_p3_fails_without_pasted_copy(tmp_path):
    files = build_returned_files(tmp_path, paste=False)
    report = check_spike(files)
    assert row(report, "P3").verdict == "FAIL"
    assert "need at least 3" in row(report, "P3").detail
    # the row must say a skipped step also produces this failure
    assert row(report, "P3").cannot_determine is not None


def test_p4_fails_on_plain_text_paste(tmp_path):
    files = build_returned_files(tmp_path)
    plain = tmp_path / "pasted.docx"
    document = Document()
    document.add_paragraph("Pasted as plain text [1]")
    document.save(plain)
    files = [f for f in files if f.name != "pasted.docx"] + [plain]
    report = check_spike(files)
    assert row(report, "P4").verdict == "FAIL"
    assert "0 tagged citation control(s)" in row(report, "P4").detail


def test_p5_fails_without_tracked_changes(tmp_path):
    files = build_returned_files(tmp_path, tracked=False)
    report = check_spike(files)
    assert row(report, "P5").verdict == "FAIL"
    assert "no tracked edits" in row(report, "P5").detail
    assert row(report, "P5").cannot_determine is not None


def test_p2_fails_when_visible_text_tampered(tmp_path):
    files = build_returned_files(tmp_path, tamper_text=True)
    report = check_spike(files)
    assert row(report, "P2").verdict == "FAIL"
    assert "'[7]'" in row(report, "P2").detail
    # and it says what the row cannot determine
    assert "cannot be attributed to step 2 alone" in row(report, "P2").cannot_determine


def test_missing_files_are_named_plainly(tmp_path):
    main = build_returned_files(tmp_path, include_step1=False, include_pasted=False, include_renamed=False)
    report = check_spike([main[0]])
    assert "step1_reopened.docx was not returned" in row(report, "P1").detail
    assert "pasted.docx was not returned" in row(report, "P4").detail
    assert "spike_renamed.docx was not returned" in row(report, "P6").detail


def test_unrecognized_file_is_reported_and_never_graded(tmp_path):
    # The old checker graded any citation-bearing file as the paste. A file
    # with a different name must be reported as unrecognized instead.
    files = build_returned_files(tmp_path, include_pasted=False)
    stray = tmp_path / "some_other_file.docx"
    document = Document()
    document.add_paragraph("Blah ")
    insert_citation(document.paragraphs[0], "C001", "[1]")
    document.save(stray)
    report = check_spike(files + [stray])
    assert any("some_other_file.docx" in line for line in report.classification)
    assert "pasted.docx was not returned" in row(report, "P4").detail
    assert "some_other_file.docx" in row(report, "P4").detail
    assert row(report, "P4").file_read is None


def test_exit_contract(tmp_path):
    good = build_returned_files(tmp_path)
    assert check_spike(good).all_passed is True
    incomplete = build_returned_files(
        tmp_path, include_step1=False, include_pasted=False, include_renamed=False
    )
    assert check_spike(incomplete).all_passed is False


def test_nonexistent_file_is_reported_not_crashing(tmp_path):
    # A typo'd path must degrade to a plain statement, never a traceback.
    files = build_returned_files(tmp_path)
    report = check_spike(files + [tmp_path / "typoed_name.docx"])
    assert any(
        "typoed_name.docx" in line and "does not exist" in line
        for line in report.classification
    )
    # the recognized files are still graded
    assert row(report, "P1").verdict == "PASS"


def test_corrupt_file_with_promised_name_is_reported_not_crashing(tmp_path):
    # Same class as the nonexistent-path bug, one level deeper: a file with
    # a promised name that is not a DOCX (wrong save format / damaged) must
    # produce a plain statement, never a traceback — and per the review
    # note's rule, "returned but unreadable" is distinct from "not returned".
    from citebind.spike import make_spike as make

    main = make(tmp_path)
    corrupt = tmp_path / "pasted.docx"
    corrupt.write_text("Saved as plain text, renamed to .docx")
    report = check_spike([main, corrupt])
    assert "was returned but is not a readable DOCX" in row(report, "P4").detail
    assert row(report, "P4").file_read is None
    # the healthy file is still graded (step1 absent -> P1 fails as not returned)
    assert row(report, "P2").verdict == "PASS"


def test_directory_with_promised_name_is_reported_not_crashing(tmp_path):
    # A directory named pasted.docx must degrade to a plain statement
    # (IsADirectoryError is an OSError, and humans pass directories).
    main = make_spike(tmp_path)
    (tmp_path / "pasted.docx").mkdir()
    report = check_spike([main, tmp_path / "pasted.docx"])
    assert "was returned but could not be opened" in row(report, "P4").detail
    assert row(report, "P2").verdict == "PASS"  # healthy file still graded


def test_promised_names_match_case_insensitively(tmp_path):
    # Windows and macOS filesystems are case-insensitive by default; a
    # returned PASTED.DOCX is the promised pasted.docx. This is still
    # name-based classification, not content inference.
    files = build_returned_files(tmp_path, include_pasted=False)
    pasted = tmp_path / "PASTED.DOCX"
    document = Document()
    document.add_paragraph("Blah ")
    insert_citation(document.paragraphs[0], "C001", "[1]")
    document.save(pasted)
    report = check_spike(files + [pasted])
    assert row(report, "P4").file_read.endswith("PASTED.DOCX")
    assert row(report, "P4").verdict == "PASS"
